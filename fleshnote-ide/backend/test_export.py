"""Manuscript export tests: every format and content mode on a fixture project
whose chapters use every construct the editor can save (export_fixture.py).

The PDF's printed layout (page size, mirrored margins) is checked by
tools/export_review.py, which needs Electron; here the print document is checked.

Run from backend/:  .venv/Scripts/python.exe -m unittest test_export -v
"""
import io
import os
import re
import shutil
import sqlite3
import tempfile
import unittest
import zipfile

import docx
from lxml import etree

import export_fixture
from export.document import build_chapter, parse_blocks
from export.pipeline import FORMATS, ExportPipeline
from export.typography import typeset

MODES = ("prose", "notes", "full")
BOOK = {"trim": "standard", "font_size": 11, "gutter": 0.625, "outer": 0.5}
LEFTOVERS = re.compile(r"\{\{|\[\[|</?x-ref|</?fn-ref|data-entity|data-time|TWIST_REF|FOOTNOTE_REF|[“”’]>")


def docx_text(data: bytes) -> str:
    return "\n".join(p.text for p in docx.Document(io.BytesIO(data)).paragraphs)


def epub_text(data: bytes) -> dict:
    z = zipfile.ZipFile(io.BytesIO(data))
    return {n: z.read(n).decode("utf-8") for n in z.namelist() if n.endswith(".xhtml")}


class ExportTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.root = tempfile.mkdtemp(prefix="fn_export_")
        cls.project = export_fixture.build(cls.root)
        cls.pipe = ExportPipeline(cls.project)
        cls.out = {}
        for mode in MODES:
            chapters, todos = cls.pipe.chapters(mode)
            cls.todos = todos
            for fmt in FORMATS:
                cls.out[(mode, fmt)] = cls.pipe.render(mode, fmt, True, BOOK, chapters)
            cls.out[(mode, "pdf-ms")] = cls.pipe.render(mode, "pdf", False, None, chapters)
            cls.out[(mode, "docx-ms")] = cls.pipe.render(mode, "docx", False, None, chapters)

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.root, ignore_errors=True)

    def texts(self, mode):
        """Every output of a mode as text (DOCX and EPUB unpacked)."""
        o = {k[1]: v for k, v in self.out.items() if k[0] == mode}
        yield "txt", o["txt"]
        yield "md", o["md"]
        yield "html", o["html"]
        yield "pdf", o["pdf"]
        yield "pdf-ms", o["pdf-ms"]
        yield "docx", docx_text(o["docx"])
        yield "docx-ms", docx_text(o["docx-ms"])
        for name, x in epub_text(o["epub"]).items():
            yield "epub " + name, x

    # ── content ──────────────────────────────────────────────────────────

    def test_deleted_chapter_is_never_exported(self):
        for mode in MODES:
            for fmt, text in self.texts(mode):
                self.assertNotIn("DELETED-CHAPTER-TEXT", text, (mode, fmt))
                self.assertNotIn("Cut Chapter", text, (mode, fmt))

    def test_no_markers_or_placeholders_leak(self):
        for mode in MODES:
            for fmt, text in self.texts(mode):
                self.assertIsNone(LEFTOVERS.search(text), (mode, fmt, LEFTOVERS.search(text)))

    def test_quotes_dashes_and_ellipses(self):
        for mode in MODES:
            for fmt, text in self.texts(mode):
                if "Light them" not in text:
                    continue
                self.assertIn("“Light them,” said", text.replace("**", ""), (mode, fmt))
                self.assertIn("’Tis late", text, (mode, fmt))
                self.assertIn("It’s always late—the", text.replace("*", "").replace("<em>", "").replace("</em>", ""), (mode, fmt))
                self.assertIn("slept…", text, (mode, fmt))
                self.assertNotIn("—-", text, (mode, fmt))

    def test_entities_unescaped_in_plain_formats_and_escaped_in_html(self):
        txt = self.out[("prose", "txt")]
        self.assertIn("All of them & quickly", txt)
        self.assertIn("Tom <3 Jerry", txt)
        self.assertIn("FIRE & ICE <PART TWO>", txt)
        self.assertNotIn("&amp;", txt)
        self.assertIn("All of them & quickly", docx_text(self.out[("prose", "docx")]))
        self.assertIn("Tom \\<3 Jerry", self.out[("prose", "md")])
        for fmt in ("html", "pdf"):
            self.assertIn("Fire &amp; Ice &lt;Part Two&gt;", self.out[("prose", fmt)], fmt)
            self.assertIn("Tom &lt;3 Jerry", self.out[("prose", fmt)], fmt)
        chapters = epub_text(self.out[("prose", "epub")])
        for name, x in chapters.items():
            etree.fromstring(x.encode("utf-8"))  # well-formed XHTML
        self.assertTrue(any("Fire &amp; Ice &lt;Part Two&gt;" in x for x in chapters.values()))

    def test_scene_breaks(self):
        txt = self.out[("prose", "txt")]
        self.assertEqual(txt.count("* * *"), 3)  # "* * *", <hr>, "---"
        self.assertEqual(self.out[("prose", "pdf")].count('class="scene"'), 3)
        self.assertEqual(self.out[("prose", "pdf-ms")].count('role="separator">#<'), 3)

    def test_inline_formatting_survives(self):
        self.assertIn("<em>lamps</em>", self.out[("prose", "html")])
        self.assertIn("<strong>cold</strong>", self.out[("prose", "html")])
        self.assertIn("*lamps*", self.out[("prose", "md")])
        self.assertIn("**cold**", self.out[("prose", "md")])
        self.assertIn("~~struck~~", self.out[("prose", "md")])
        self.assertIn("[link](https://example.org)", self.out[("prose", "md")])
        runs = [r for p in docx.Document(io.BytesIO(self.out[("prose", "docx")])).paragraphs for r in p.runs]
        self.assertTrue(any(r.text == "lamps" and r.italic for r in runs))
        self.assertTrue(any(r.text == "cold" and r.bold for r in runs))
        self.assertTrue(any(r.text == "struck" and r.font.strike for r in runs))

    def test_line_breaks_stay_inside_the_paragraph(self):
        self.assertIn("The year before,\nthe wall fell.", self.out[("prose", "txt")])
        self.assertIn("The year before,<br/>the wall fell.", self.out[("prose", "html")])
        self.assertIn("The year before,  \nthe wall fell.", self.out[("prose", "md")])

    def test_opening_heading_is_not_repeated(self):
        txt = self.out[("prose", "txt")]
        self.assertEqual(txt.count("Chapter 1"), 1)
        self.assertIn("Chapter 1\nTHE LAMPS", txt)
        self.assertIn('<p class="chapter-label">Chapter 1</p>', self.out[("prose", "pdf")])

    def test_quick_note_keeps_the_text_it_marks(self):
        for mode in MODES:
            for fmt, text in self.texts(mode):
                if "Light them" in text:
                    self.assertIn("tide", text, (mode, fmt))
                self.assertNotIn("QUICKNOTE-SECRET", text, (mode, fmt))

    def test_annotations_become_footnotes_only_with_notes(self):
        note = export_fixture.ANNOTATION_TEXT
        for fmt, text in self.texts("prose"):
            self.assertNotIn(note, text, fmt)
        for mode in ("notes", "full"):
            self.assertIn("harbor[1] slept", self.out[(mode, "txt")])
            self.assertIn("[1] " + note, self.out[(mode, "txt")])
            self.assertIn("harbor[^1] slept", self.out[(mode, "md")])
            self.assertIn("[^1]: " + note, self.out[(mode, "md")])
            self.assertIn('id="ch1fn1"', self.out[(mode, "pdf")])
            self.assertIn(note, docx_text(self.out[(mode, "docx")]))

    def test_full_mode_shows_links_and_twists(self):
        html = self.out[("full", "html")]
        self.assertIn('<span class="entity char">Mara</span>', html)
        self.assertIn('title="%s"' % export_fixture.LORE_DESC, html)
        self.assertIn('<span class="twist">It was him</span>', html)
        self.assertIn("**Mara**", self.out[("full", "md")])
        self.assertNotIn('class="entity', self.out[("prose", "html")])

    def test_todo_notes_are_removed_and_counted(self):
        self.assertEqual(self.todos, 1)
        for mode in MODES:
            for fmt, text in self.texts(mode):
                self.assertNotIn("#TODO", text, (mode, fmt))

    def test_chapter_selection(self):
        chapters, _ = self.pipe.chapters("prose", ["ch-2"])
        self.assertEqual([c.title for c in chapters], ["Fire & Ice <Part Two>"])
        chapters, _ = self.pipe.chapters("prose", ["ch-3"])  # deleted
        self.assertEqual(chapters, [])

    # ── layout ───────────────────────────────────────────────────────────

    def test_print_book_page_setup(self):
        html = self.out[("prose", "pdf")]
        self.assertIn("size: 5.0in 8.0in", html)
        self.assertIn("@page :left { margin-left: 0.5in; margin-right: 0.625in; }", html)
        self.assertIn("@page :right { margin-left: 0.625in; margin-right: 0.5in; }", html)
        self.assertIn("counter(page)", html)
        self.assertIn("­", html)  # soft hyphens for justified text

    def test_print_manuscript_page_setup(self):
        html = self.out[("prose", "pdf-ms")]
        self.assertIn("size: 8.5in 11in; margin: 1in", html)
        self.assertIn('content: "Quill / EXPORT FIXTURE / " counter(page)', html)
        self.assertIn("line-height: 2", html)

    def test_docx_book_and_manuscript(self):
        book = docx.Document(io.BytesIO(self.out[("prose", "docx")]))
        s = book.sections[0]
        self.assertAlmostEqual(s.page_width.inches, 5.0, places=2)
        self.assertAlmostEqual(s.page_height.inches, 8.0, places=2)
        self.assertAlmostEqual(s.left_margin.inches, 0.625, places=2)
        settings = book.settings.element.xml
        self.assertIn("w:mirrorMargins", settings)
        self.assertIn("w:doNotExpandShiftReturn", settings)
        self.assertLess(settings.index("w:mirrorMargins"), settings.index("w:defaultTabStop"))
        ms = docx.Document(io.BytesIO(self.out[("prose", "docx-ms")]))
        s = ms.sections[0]
        self.assertAlmostEqual(s.page_width.inches, 8.5, places=2)
        self.assertAlmostEqual(s.left_margin.inches, 1.0, places=2)
        self.assertNotIn("w:mirrorMargins", ms.settings.element.xml)
        self.assertIn("EXPORT FIXTURE", s.header.paragraphs[0].text)

    def test_epub_language_follows_the_manuscript(self):
        conn = sqlite3.connect(os.path.join(self.project, "fleshnote.db"))
        conn.execute("UPDATE project_config SET config_value = 'hu' WHERE config_key = 'story_language'")
        conn.commit()
        conn.close()
        try:
            pipe = ExportPipeline(self.project)
            chapters, _ = pipe.chapters("prose")
            data = pipe.render("prose", "epub", chapters=chapters)
            opf = [zipfile.ZipFile(io.BytesIO(data)).read(n).decode() for n in
                   zipfile.ZipFile(io.BytesIO(data)).namelist() if n.endswith(".opf")][0]
            self.assertIn("<dc:language>hu</dc:language>", opf)
            self.assertIn("„Light them,”", pipe.render("prose", "txt", chapters=chapters))
        finally:
            conn = sqlite3.connect(os.path.join(self.project, "fleshnote.db"))
            conn.execute("UPDATE project_config SET config_value = 'en' WHERE config_key = 'story_language'")
            conn.commit()
            conn.close()

    def test_exports_never_overwrite_each_other(self):
        paths = [self.pipe.run("prose", "txt")[0] for _ in range(3)]
        self.assertEqual(len(set(paths)), 3)
        self.assertTrue(all(os.path.exists(p) for p in paths))
        pdf_path, _, print_html = self.pipe.run("prose", "pdf", True, BOOK)
        self.assertTrue(pdf_path.endswith(".pdf"))
        self.assertFalse(os.path.exists(pdf_path))  # the app prints it
        self.assertIn("@page", print_html)


class TypographyTests(unittest.TestCase):
    def test_quotes_open_after_markup_boundaries(self):
        self.assertEqual(typeset('"Hi," she said.'), "“Hi,” she said.")
        self.assertEqual(typeset('" she said.', prev="i"), "” she said.")
        self.assertEqual(typeset('"word"', prev="("), "“word”")

    def test_apostrophes(self):
        self.assertEqual(typeset("'Tis the night of the '90s, rock 'n' roll."),
                         "’Tis the night of the ’90s, rock ’n’ roll.")
        self.assertEqual(typeset("It's the boys' toys."), "It’s the boys’ toys.")
        self.assertEqual(typeset("'Run,' he said."), "‘Run,’ he said.")

    def test_language_quotes(self):
        self.assertEqual(typeset('"Szia" mondta.', lang="hu"), "„Szia” mondta.")
        self.assertEqual(typeset('"Cześć" - rzekł.', lang="pl"), "„Cześć” - rzekł.")

    def test_dashes_ellipses_spaces(self):
        self.assertEqual(typeset("late--and...  then---gone"), "late—and… then—gone")


class ParserTests(unittest.TestCase):
    def test_plain_text_chapters(self):
        blocks = parse_blocks("# Chapter 1\n\nFirst line\nsame paragraph.\n\nSecond & last.")
        self.assertEqual([b.kind for b in blocks], ["h", "p", "p"])
        self.assertEqual(blocks[1].plain, "First line\nsame paragraph.")
        self.assertEqual(blocks[2].plain, "Second & last.")

    def test_scene_break_variants(self):
        for marker in ("* * *", "***", "#", "---", "⁂", "~"):
            blocks = parse_blocks("<p>a</p><p>%s</p><p>b</p>" % marker)
            self.assertEqual([b.kind for b in blocks], ["p", "scene", "p"], marker)

    def test_title_and_heading(self):
        same = build_chapter("Chapter 1", "<h1>Chapter 1</h1><p>x</p>", [])
        self.assertEqual((same.title, same.label), ("Chapter 1", ""))
        label = build_chapter("The Lamps", "<h1>Chapter 1</h1><p>x</p>", [])
        self.assertEqual((label.title, label.label), ("The Lamps", "Chapter 1"))
        longer = build_chapter("Chapter 1", "<h1>Chapter 1: The Lamps</h1><p>x</p>", [])
        self.assertEqual((longer.title, longer.label), ("Chapter 1: The Lamps", ""))
        untitled = build_chapter("", "<h2>Prologue</h2><p>x</p>", [])
        self.assertEqual(untitled.title, "Prologue")

    def test_mention_and_unknown_tags_keep_their_text(self):
        blocks = parse_blocks('<p>Ask <span data-type="mention" data-id="1">@Mara</span> <mark>now</mark>.</p>')
        self.assertEqual(blocks[0].plain, "Ask @Mara now.")


if __name__ == "__main__":
    unittest.main()
