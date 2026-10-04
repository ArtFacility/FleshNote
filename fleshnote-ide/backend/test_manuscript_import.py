import os
import shutil
import sqlite3
import sys
import tempfile
import unittest
from contextlib import closing

backend_dir = os.path.dirname(os.path.abspath(__file__))
if backend_dir not in sys.path:
    sys.path.insert(0, backend_dir)

import db_setup
from manuscript_import import (
    SCENE_BREAK,
    inline_to_html,
    is_chapter_heading,
    paragraphs_to_html,
    preview_files,
    split_blocks,
    text_to_blocks,
)


def split_text(text, **kw):
    return split_blocks(text_to_blocks(text), **kw)


def filler(words):
    sentence = "The tide came in slowly over the black stones of the harbor. "
    n = len(sentence.split())
    return (sentence * (words // n + 1)).strip()


class TestHeadings(unittest.TestCase):
    def test_chapter_headings(self):
        for line in [
            "Chapter 1", "CHAPTER IV", "Chapter One", "Chapter twenty-one", "Chapter 3: The Harbor",
            "Chapter 3 The Sea", "Part II", "Act 1", "Prologue", "Epilogue: After",
            "1. fejezet", "Első fejezet", "Rozdział 3", "Rozdział pierwszy", "Część druga",
            "الفصل الأول",
        ]:
            self.assertTrue(is_chapter_heading(line), line)

    def test_prose_is_not_a_heading(self):
        for line in [
            "Actually, she had never seen the sea.", "Partly it was fear.",
            "Booked passage was expensive.", "Part of me wanted to go.", "Act civil.",
            "Chapters of my life", "Book of Shadows", "Prologues are boring.",
            "Chapter 3 begins where the last one ended so long ago.",
        ]:
            self.assertFalse(is_chapter_heading(line), line)


class TestSplitting(unittest.TestCase):
    def test_prose_starting_with_keywords_stays_in_the_chapter(self):
        text = (
            "Chapter 1: The Harbor\n\nActually, she had never seen the sea.\n\n"
            "Partly it was fear.\n\nBooked passage was expensive.\n\n"
            "Chapter 2\n\nThe ship left at dawn."
        )
        chunks = split_text(text)
        self.assertEqual([c["title"] for c in chunks], ["Chapter 1: The Harbor", "Chapter 2"])
        self.assertEqual(len(chunks[0]["paragraphs"]), 3)

    def test_localized_headings_split(self):
        text = "1. fejezet\n\nHajnalban indultak.\n\nRozdział 2\n\nStatek odpłynął.\n\nالفصل الأول\n\nنص."
        self.assertEqual(len(split_text(text)), 3)

    def test_scene_break_stays_inside_a_chapter(self):
        text = "Chapter 1\n\nFirst scene.\n\n***\n\nSecond scene.\n\nChapter 2\n\nMore."
        chunks = split_text(text)
        self.assertEqual(len(chunks), 2)
        self.assertEqual(chunks[0]["paragraphs"], ["First scene.", SCENE_BREAK, "Second scene."])

    def test_table_of_contents_and_front_matter_are_flagged(self):
        text = (
            "The Salt Crown\n\nby A. Writer\n\n"
            "Chapter 1 ........ 3\n\nChapter 2 ........ 9\n\nChapter 3 ........ 14\n\n"
            f"Chapter 1\n\n{filler(60)}\n\nChapter 2\n\n{filler(60)}\n\nChapter 3\n\n{filler(60)}"
        )
        chunks = split_text(text)
        self.assertEqual(chunks[0]["flag"], "front")
        self.assertEqual([c["flag"] for c in chunks[1:4]], ["empty", "empty", "empty"])
        self.assertEqual([c["flag"] for c in chunks[4:]], [None, None, None])

    def test_one_line_per_paragraph_files(self):
        lines = ["Chapter 1"] + [f"Line {i} of the story goes here." for i in range(15)]
        lines += ["Chapter 2"] + [f"More story {i}." for i in range(10)]
        chunks = split_text("\n".join(lines))
        self.assertEqual(len(chunks), 2)
        self.assertEqual(len(chunks[0]["paragraphs"]), 15)

    def test_hard_wrapped_lines_are_rejoined(self):
        para = (
            "It was a bright cold day in April, and the clocks were striking\n"
            "thirteen. Winston Smith, his chin nuzzled into his breast in an\n"
            "effort to escape the vile wind, slipped quickly through the glass"
        )
        text = "Chapter 1\n\n" + "\n\n".join([para] * 6)
        chunks = split_text(text)
        self.assertNotIn("\n", chunks[0]["paragraphs"][0])

    def test_heading_glued_to_its_first_paragraph(self):
        chunks = split_text("Chapter 1\nIt was late.\n\nChapter 2\nIt was early.")
        self.assertEqual([c["title"] for c in chunks], ["Chapter 1", "Chapter 2"])
        self.assertEqual(chunks[1]["paragraphs"], ["It was early."])

    def test_markdown_levels(self):
        text = "# The Book\n\n## One\n\nText a.\n\n## Two\n\nText b."
        chunks = split_text(text)
        self.assertEqual([c["title"] for c in chunks], ["Front matter", "One", "Two"])

    def test_bare_numbers_need_chapter_spacing(self):
        chapters = "\n\n".join(f"{i}\n\n{filler(900)}" for i in range(1, 5))
        self.assertEqual(len(split_text(chapters)), 4)
        pages = "\n\n".join(f"{i}\n\n{filler(250)}" for i in range(1, 9))
        self.assertEqual(len(split_text(pages)), 1)

    def test_parts_with_numbered_chapters(self):
        part = lambda name: f"{name}\n\nI.\n\n{filler(900)}\n\nII.\n\n{filler(900)}"
        chunks = split_text(part("ELSŐ RÉSZ. A KEZDET") + "\n\n" + part("MÁSODIK RÉSZ. A VÉG"))
        self.assertEqual([c["title"] for c in chunks], ["I.", "II.", "I.", "II."])
        self.assertEqual(chunks[0]["paragraphs"][0], "### ELSŐ RÉSZ. A KEZDET")
        self.assertEqual(chunks[2]["paragraphs"][0], "### MÁSODIK RÉSZ. A VÉG")

    def test_numeral_dot_title_headings(self):
        text = "\n\n".join(f"{n}. Jak wygląda firma\n\n{filler(900)}" for n in ["I", "II", "III"])
        self.assertEqual([c["title"] for c in split_text(text)],
                         ["I. Jak wygląda firma", "II. Jak wygląda firma", "III. Jak wygląda firma"])

    def test_numbered_list_in_prose_is_not_chapters(self):
        text = "Chapter 1\n\nShe made a list.\n\n1. Bread\n\n2. Milk\n\n3. Eggs\n\nThen she left."
        self.assertEqual(len(split_text(text)), 1)

    def test_ebook_licence_becomes_back_matter(self):
        text = (f"Chapter 1\n\n{filler(60)}\n\nChapter 2\n\n{filler(60)}\n\n"
                "*** END OF THE PROJECT GUTENBERG EBOOK THE BOOK ***\n\nUpdated editions will replace the previous one.")
        chunks = split_text(text)
        self.assertEqual([c["flag"] for c in chunks], [None, None, "back"])
        self.assertNotIn("Updated", chunks[1]["content"])

    def test_no_headings_and_short_scenes_stays_one_chapter(self):
        text = "A scene.\n\n* * *\n\nAnother scene."
        self.assertEqual(len(split_text(text)), 1)

    def test_no_headings_chapter_sized_breaks_split(self):
        text = f"{filler(4500)}\n\n***\n\n{filler(4500)}"
        self.assertEqual(len(split_text(text)), 2)

    def test_one_file_per_chapter_mode(self):
        chunks = split_text("Chapter 7: Ash\n\nText.\n\n***\n\nMore.", fallback_title="ch07",
                            one_chapter_unless_headed=True)
        self.assertEqual(len(chunks), 1)
        self.assertEqual(chunks[0]["title"], "Chapter 7: Ash")
        chunks = split_text("Just text.", fallback_title="ch07", one_chapter_unless_headed=True)
        self.assertEqual(chunks[0]["title"], "ch07")


class TestHtml(unittest.TestCase):
    def test_escaping_and_emphasis(self):
        self.assertEqual(inline_to_html("a < b & *c* **d** _e_"),
                         "a &lt; b &amp; <em>c</em> <strong>d</strong> <em>e</em>")
        self.assertEqual(inline_to_html(r"5 \* 3 and snake_case"), "5 * 3 and snake_case")
        self.assertEqual(inline_to_html("**_both_**"), "<strong><em>both</em></strong>")

    def test_blocks(self):
        html = paragraphs_to_html(["One", SCENE_BREAK, "### Sub", "Two\nlines"])
        self.assertEqual(html, "<p>One</p><p>* * *</p><h3>Sub</h3><p>Two<br>lines</p>")


class TestFiles(unittest.TestCase):
    def setUp(self):
        self.dir = tempfile.mkdtemp(prefix="fn_import_")

    def tearDown(self):
        shutil.rmtree(self.dir, ignore_errors=True)

    def write(self, name, text, encoding="utf-8"):
        path = os.path.join(self.dir, name)
        with open(path, "w", encoding=encoding) as f:
            f.write(text)
        return path

    def test_folder_of_chapter_files_in_natural_order(self):
        self.write("ch10.txt", "Ten.")
        self.write("ch2.txt", "Two.")
        self.write("ch1.txt", "One.")
        self.write("notes.pdf", "ignored")
        result = preview_files([self.dir])
        self.assertEqual([s["title"] for s in result["splits"]], ["ch1", "ch2", "ch10"])

    def test_files_added_one_at_a_time_keep_their_names(self):
        a = self.write("03 - The Harbor.txt", "Text of three.\n\n\n\n\n\nStill three.")
        b = self.write("04 - The Sea.txt", "Text of four.")
        titles = [preview_files([p])["splits"][0]["title"] for p in (a, b)]
        self.assertEqual(titles, ["03 - The Harbor", "04 - The Sea"])
        self.assertEqual(preview_files([a])["total_chapters"], 1)

    def test_download_copy_suffix_is_dropped_from_titles(self):
        self.write("Chapter 3 (1).txt", "Text.")
        self.write("Chapter 4 (1).txt", "Text.")
        result = preview_files([self.dir])
        self.assertEqual([s["title"] for s in result["splits"]], ["Chapter 3", "Chapter 4"])

    def test_bad_file_does_not_sink_the_batch(self):
        good = self.write("a.txt", "Text.")
        bad = self.write("b.odt", "x")
        result = preview_files([good, bad])
        self.assertEqual(result["total_chapters"], 1)
        self.assertIn("error", result["files"][1])

    def test_windows_1250_text(self):
        path = self.write("pl.txt", "Rozdział 1\n\nZażółć gęślą jaźń.", encoding="cp1250")
        result = preview_files([path])
        self.assertEqual(result["splits"][0]["paragraphs"], ["Zażółć gęślą jaźń."])

    def test_docx_styles_and_emphasis(self):
        import docx
        d = docx.Document()
        d.add_paragraph("The Salt Crown", style="Title")
        d.add_paragraph("Chapter One", style="Heading 1")
        p = d.add_paragraph("She was ")
        p.add_run("very").italic = True
        p.add_run(" tired, 2*3.")
        d.add_paragraph("Chapter Two", style="Heading 1")
        d.add_paragraph("Done.")
        path = os.path.join(self.dir, "book.docx")
        d.save(path)
        splits = preview_files([path])["splits"]
        self.assertEqual([s["title"] for s in splits], ["Front matter", "Chapter One", "Chapter Two"])
        self.assertEqual(paragraphs_to_html(splits[1]["paragraphs"]), "<p>She was <em>very</em> tired, 2*3.</p>")


class TestTrickyFiles(unittest.TestCase):
    def setUp(self):
        self.dir = tempfile.mkdtemp(prefix="fn_tricky_")

    def tearDown(self):
        shutil.rmtree(self.dir, ignore_errors=True)

    def test_docx_hyperlinks_and_tracked_insertions_keep_their_words(self):
        import docx
        from docx.oxml.ns import qn
        from docx.oxml import OxmlElement

        d = docx.Document()
        p = d.add_paragraph("See ")
        link = OxmlElement("w:hyperlink")
        run = OxmlElement("w:r")
        text = OxmlElement("w:t")
        text.text = "the map"
        run.append(text)
        link.append(run)
        p._p.append(link)
        ins = OxmlElement("w:ins")
        ins.set(qn("w:id"), "1")
        ins.set(qn("w:author"), "Editor")
        run2 = OxmlElement("w:r")
        text2 = OxmlElement("w:t")
        text2.text = " before dawn"
        text2.set(qn("xml:space"), "preserve")
        run2.append(text2)
        ins.append(run2)
        p._p.append(ins)
        path = os.path.join(self.dir, "linked.docx")
        d.save(path)
        para = preview_files([path])["splits"][0]["paragraphs"][0]
        self.assertEqual(para, "See the map before dawn")

    def test_utf16_text_export(self):
        path = os.path.join(self.dir, "unicode.txt")
        with open(path, "wb") as f:
            f.write("Fejezet nélkül.\n\nÁrvíztűrő tükörfúrógép.".encode("utf-16"))
        self.assertEqual(preview_files([path])["splits"][0]["paragraphs"],
                         ["Fejezet nélkül.", "Árvíztűrő tükörfúrógép."])


class TestConfirm(unittest.TestCase):
    def setUp(self):
        self.dir = tempfile.mkdtemp(prefix="fn_confirm_")
        self.project = os.path.join(self.dir, "proj")
        os.makedirs(os.path.join(self.project, "md"))
        db_setup.generate_project_db(self.project, {"project_name": "T"})
        from framework_presets import _create_single_default_chapter  # noqa: F401  (created by init)

    def tearDown(self):
        shutil.rmtree(self.dir, ignore_errors=True)

    def db(self):
        return closing(sqlite3.connect(os.path.join(self.project, "fleshnote.db")))

    def live(self):
        with self.db() as c:
            return c.execute(
                "SELECT chapter_number, title FROM chapters WHERE deleted = 0 ORDER BY chapter_number"
            ).fetchall()

    def confirm(self, titles, **kw):
        from routes.imports import confirm_splits, ConfirmSplitsRequest
        splits = [{"title": t, "paragraphs": [f"Text of {t}."]} for t in titles]
        return confirm_splits(ConfirmSplitsRequest(project_path=self.project, splits=splits, **kw))

    def test_replaces_the_new_project_placeholder(self):
        self.confirm(["A", "B"], replace_placeholder=True)
        self.assertEqual(self.live(), [(1, "A"), (2, "B")])

    def test_keeps_chapter_one_unless_asked(self):
        self.confirm(["A"])
        self.assertEqual(self.live(), [(1, "Chapter 1"), (2, "A")])

    def test_keeps_a_chapter_one_with_writing_in_it(self):
        with self.db() as c:
            c.execute("UPDATE chapters SET word_count = 12, status = 'writing'")
            c.commit()
        self.confirm(["A"], replace_placeholder=True)
        self.assertEqual(self.live(), [(1, "Chapter 1"), (2, "A")])

    def test_insert_after_shifts_later_chapters_and_logs_them(self):
        self.confirm(["A", "B", "C"], replace_placeholder=True)
        self.confirm(["X", "Y"], insert_after=1)
        self.assertEqual([t for _, t in self.live()], ["A", "X", "Y", "B", "C"])
        with self.db() as c:
            rows = c.execute(
                "SELECT COUNT(*) FROM change_log WHERE table_name = 'chapters' AND column_name = 'chapter_number'"
            ).fetchone()[0]
        self.assertGreaterEqual(rows, 5 + 2)  # every created chapter + the two shifted

    def test_files_are_written_as_html(self):
        result = self.confirm(["A"], replace_placeholder=True)
        name = result["chapters"][0]["md_filename"]
        with open(os.path.join(self.project, "md", name), encoding="utf-8") as f:
            self.assertEqual(f.read(), "<p>Text of A.</p>")

    def test_deleting_a_middle_chapter_renumbers(self):
        from routes.chapters import delete_chapter, ChapterDelete
        self.confirm(["A", "B", "C"], replace_placeholder=True)
        with self.db() as c:
            b_id = c.execute("SELECT id FROM chapters WHERE title = 'B'").fetchone()[0]
        delete_chapter(ChapterDelete(project_path=self.project, chapter_id=b_id))
        self.assertEqual(self.live(), [(1, "A"), (2, "C")])
        self.confirm(["D"])
        self.assertEqual(self.live(), [(1, "A"), (2, "C"), (3, "D")])


class TestGenreDefaults(unittest.TestCase):
    def test_missing_settings_come_from_the_genre(self):
        filled = db_setup.with_genre_defaults({"genre": "fantasy", "track_groups": False})
        self.assertTrue(filled["track_species"])
        self.assertFalse(filled["track_groups"])
        self.assertIn("artifact", filled["lore_categories"])


if __name__ == "__main__":
    unittest.main()
